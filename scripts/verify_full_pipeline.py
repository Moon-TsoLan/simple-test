# -*- coding: utf-8 -*-
"""Verify a full task-1/task-2 run through the application API.

This is an engineering acceptance test, not an accuracy evaluation.  It
audits source coordinates, exercises every public query family, materializes
full CSV/XLSX exports, and records response times in a machine-readable report.
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import sqlite3
import sys
import time
from pathlib import Path
from typing import Any

from fastapi.testclient import TestClient
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src.api.app import create_app  # noqa: E402


def _json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def _valid_ref(ref: dict[str, Any], blocks: dict[str, dict[str, Any]]) -> tuple[bool, str | None]:
    block_id = str(ref.get("block_id") or "")
    if block_id not in blocks:
        return False, "unknown_block"
    row = ref.get("block_row")
    if row is None:
        return True, None
    rows = blocks[block_id].get("rows") or []
    if not isinstance(row, int) or not 1 <= row <= len(rows):
        return False, "invalid_row"
    return True, None


def audit_task1(task1_dir: Path, blocks_dir: Path) -> dict[str, Any]:
    counts = {
        "notices": 0, "items": 0, "items_with_primary_reference": 0,
        "items_with_evidence_refs": 0, "valid_primary_references": 0,
        "valid_evidence_references": 0, "invalid_primary_references": 0,
        "invalid_evidence_references": 0,
    }
    failures: list[dict[str, Any]] = []
    for path in sorted(task1_dir.glob("*.json")):
        notice = _json(path)
        notice_id = str(notice.get("notice_id") or path.stem)
        block_path = blocks_dir / f"{notice_id}.json"
        blocks = {
            str(value.get("block_id")): value
            for value in (_json(block_path).get("blocks") if block_path.is_file() else []) or []
        }
        counts["notices"] += 1
        for item_index, item in enumerate(notice.get("items") or [], 1):
            counts["items"] += 1
            primary = {
                "block_id": item.get("source_block_id"),
                "block_row": item.get("source_block_row"),
            }
            if primary["block_id"]:
                counts["items_with_primary_reference"] += 1
                valid, reason = _valid_ref(primary, blocks)
                counts["valid_primary_references" if valid else "invalid_primary_references"] += 1
                if not valid and len(failures) < 50:
                    failures.append({"notice_id": notice_id, "item_no": item_index, "kind": "primary", "reason": reason, **primary})
            refs = item.get("evidence_refs") or []
            if refs:
                counts["items_with_evidence_refs"] += 1
            for ref in refs:
                valid, reason = _valid_ref(ref, blocks)
                counts["valid_evidence_references" if valid else "invalid_evidence_references"] += 1
                if not valid and len(failures) < 50:
                    failures.append({"notice_id": notice_id, "item_no": item_index, "kind": "evidence", "reason": reason, **ref})
    return {**counts, "sample_failures": failures}


def audit_task2(relation_dir: Path, blocks_dir: Path) -> dict[str, Any]:
    counts = {
        "notices": 0, "success_notices": 0, "packages": 0, "bids": 0,
        "bids_with_any_reference": 0, "bids_without_reference": 0,
        "valid_references": 0, "invalid_references": 0,
        "projects_with_procurement_unit": 0, "projects_with_agency": 0,
        "projects_with_bids": 0, "projects_with_winner": 0,
        "bids_with_amount": 0, "bids_with_score": 0, "bids_with_rank": 0,
        "bids_with_known_qualification": 0, "bids_with_known_compliance": 0,
    }
    failures: list[dict[str, Any]] = []
    methods: dict[str, int] = {}
    for path in sorted(relation_dir.glob("*.json")):
        outcome = _json(path)
        notice_id = str(outcome.get("notice_id") or path.stem)
        project = outcome.get("project") or {}
        block_path = blocks_dir / f"{notice_id}.json"
        blocks = {
            str(value.get("block_id")): value
            for value in (_json(block_path).get("blocks") if block_path.is_file() else []) or []
        }
        counts["notices"] += 1
        if outcome.get("status") == "success":
            counts["success_notices"] += 1
        method = str(outcome.get("method") or "unknown")
        methods[method] = methods.get(method, 0) + 1
        if project.get("procurement_unit"):
            counts["projects_with_procurement_unit"] += 1
        if project.get("agency"):
            counts["projects_with_agency"] += 1
        packages = project.get("packages") or []
        counts["packages"] += len(packages)
        project_bids = 0
        project_winners = 0
        for package_index, package in enumerate(packages):
            for bid_index, bid in enumerate(package.get("bidders") or []):
                counts["bids"] += 1
                project_bids += 1
                if bid.get("result") in {"中标", "成交"}:
                    project_winners += 1
                for field, key in (
                    ("amount_yuan", "bids_with_amount"), ("score", "bids_with_score"),
                    ("rank", "bids_with_rank"),
                ):
                    if bid.get(field) is not None:
                        counts[key] += 1
                if bid.get("qualification") not in {None, "未知"}:
                    counts["bids_with_known_qualification"] += 1
                if bid.get("compliance") not in {None, "未知"}:
                    counts["bids_with_known_compliance"] += 1
                refs = list(bid.get("evidence_refs") or [])
                if bid.get("evidence_block_id"):
                    legacy = {"block_id": bid.get("evidence_block_id"), "block_row": bid.get("evidence_block_row")}
                    if legacy not in refs:
                        refs.append(legacy)
                if refs:
                    counts["bids_with_any_reference"] += 1
                else:
                    counts["bids_without_reference"] += 1
                for ref in refs:
                    valid, reason = _valid_ref(ref, blocks)
                    counts["valid_references" if valid else "invalid_references"] += 1
                    if not valid and len(failures) < 50:
                        failures.append({
                            "notice_id": notice_id, "package_index": package_index,
                            "bid_index": bid_index, "org": (bid.get("org") or {}).get("name"),
                            "reason": reason, **ref,
                        })
        if project_bids:
            counts["projects_with_bids"] += 1
        if project_winners:
            counts["projects_with_winner"] += 1
    return {**counts, "methods": methods, "sample_failures": failures}


def _one(connection: sqlite3.Connection, sql: str) -> tuple[Any, ...] | None:
    row = connection.execute(sql).fetchone()
    return tuple(row) if row else None


def _timed_get(client: TestClient, path: str, params: dict[str, Any] | None = None) -> tuple[Any, float]:
    started = time.perf_counter()
    response = client.get(path, params=params)
    elapsed_ms = (time.perf_counter() - started) * 1000
    response.raise_for_status()
    return response, round(elapsed_ms, 3)


def verify_api(db_path: Path, blocks_dir: Path, export_dir: Path) -> dict[str, Any]:
    export_dir.mkdir(parents=True, exist_ok=True)
    samples: dict[str, tuple[Any, ...] | None] = {}
    with sqlite3.connect(db_path) as connection:
        samples["win_unit"] = _one(connection, "SELECT unit_id FROM unit_win_stats ORDER BY win_count DESC LIMIT 1")
        samples["bid_unit"] = _one(connection, "SELECT unit_id FROM unit_bidder_stats ORDER BY bid_count DESC LIMIT 1")
        samples["pair_unit"] = _one(connection, "SELECT unit_id FROM unit_cobid_pairs ORDER BY project_count DESC LIMIT 1")
        samples["cobid_supplier"] = _one(connection, "SELECT supplier_id FROM supplier_cobid ORDER BY project_count DESC LIMIT 1")
        samples["common_pair"] = _one(connection, "SELECT supplier_a,supplier_b FROM supplier_common_units ORDER BY joint_win_count DESC LIMIT 1")
        samples["joint_pair"] = _one(connection, "SELECT supplier_a,supplier_b FROM supplier_joint_projects LIMIT 1")
        samples["graph_project"] = _one(connection, "SELECT p.project_id FROM packages p JOIN bids b ON b.package_id=p.package_id GROUP BY p.project_id ORDER BY COUNT(*) DESC LIMIT 1")

    app = create_app(db_path, frontend_dist=export_dir / "no-frontend", upload_dir=export_dir / "uploads", block_dir=blocks_dir)
    report: dict[str, Any] = {"timings_ms": {}, "scenarios": {}, "failures": []}
    with TestClient(app) as client:
        response, elapsed = _timed_get(client, "/api/status")
        report["timings_ms"]["status"] = elapsed
        report["status"] = response.json()

        response, elapsed = _timed_get(client, "/api/task1/items", {"page": 1, "page_size": 100})
        report["timings_ms"]["task1_search"] = elapsed
        page = response.json()
        report["task1_search"] = {"total": page["total"], "returned": len(page["items"])}
        if page["items"]:
            entity_id = page["items"][0]["entity_id"]
            detail_response, detail_elapsed = _timed_get(client, f"/api/task1/items/{entity_id}")
            detail = detail_response.json()
            report["timings_ms"]["task1_detail"] = detail_elapsed
            report["task1_detail"] = {
                "entity_id": entity_id,
                "has_evidence_excerpt": bool(detail.get("evidence_excerpt")),
                "source_block_id": detail.get("source_block_id"),
            }

        csv_response, csv_elapsed = _timed_get(client, "/api/task1/export.csv")
        csv_path = export_dir / "task1_entities.csv"
        csv_path.write_bytes(csv_response.content)
        csv_count = sum(1 for _ in csv.DictReader(io.StringIO(csv_response.content.decode("utf-8-sig"))))
        report["timings_ms"]["task1_export_csv"] = csv_elapsed

        xlsx_response, xlsx_elapsed = _timed_get(client, "/api/task1/export.xlsx")
        xlsx_path = export_dir / "task1_entities.xlsx"
        xlsx_path.write_bytes(xlsx_response.content)
        workbook = load_workbook(io.BytesIO(xlsx_response.content), read_only=True, data_only=True)
        xlsx_count = max(0, workbook.active.max_row - 1)
        workbook.close()
        report["timings_ms"]["task1_export_xlsx"] = xlsx_elapsed
        report["exports"] = {
            "csv": {"path": str(csv_path), "rows": csv_count, "bytes": len(csv_response.content)},
            "xlsx": {"path": str(xlsx_path), "rows": xlsx_count, "bytes": len(xlsx_response.content)},
        }

        endpoints: list[tuple[str, str, dict[str, Any] | None]] = []
        if samples["win_unit"]:
            endpoints.append(("unit_win_suppliers", f"/api/task2/units/{samples['win_unit'][0]}/win-suppliers", None))
        if samples["bid_unit"]:
            endpoints.append(("unit_top_bidders", f"/api/task2/units/{samples['bid_unit'][0]}/top-bidders", {"top": 20}))
        if samples["pair_unit"]:
            endpoints.append(("unit_cobid_pairs", f"/api/task2/units/{samples['pair_unit'][0]}/co-bid-pairs", None))
        if samples["cobid_supplier"]:
            endpoints.append(("supplier_cobidders", f"/api/task2/suppliers/{samples['cobid_supplier'][0]}/co-bidders", {"top": 20}))
        if samples["common_pair"]:
            endpoints.append(("supplier_common_units", "/api/task2/suppliers/common-units", {"ids": ",".join(samples["common_pair"])}))
        if samples["joint_pair"]:
            endpoints.append(("supplier_joint_projects", "/api/task2/suppliers/joint-projects", {"ids": ",".join(samples["joint_pair"])}))
        if samples["graph_project"]:
            endpoints.append(("project_subgraph", "/api/task2/graph/subgraph", {"project_id": samples["graph_project"][0]}))

        for name, endpoint, params in endpoints:
            response, elapsed = _timed_get(client, endpoint, params)
            value = response.json()
            count = len(value.get("nodes") or []) if isinstance(value, dict) else len(value)
            report["timings_ms"][name] = elapsed
            report["scenarios"][name] = {
                "result_count": count,
                "non_empty": count > 0,
                "sample": value if count <= 3 else (value[:3] if isinstance(value, list) else {
                    "nodes": (value.get("nodes") or [])[:3], "edges": (value.get("edges") or [])[:3],
                }),
            }

    expected = report["status"]["task1"]["entities"]
    if report["exports"]["csv"]["rows"] != expected or report["exports"]["xlsx"]["rows"] != expected:
        report["failures"].append("task1 export row count differs from database count")
    for name, value in report["scenarios"].items():
        if not value["non_empty"]:
            report["failures"].append(f"{name} returned no data for a database-derived sample")
    report["passed"] = not report["failures"]
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--db", type=Path, required=True)
    parser.add_argument("--task1-dir", type=Path, required=True)
    parser.add_argument("--relation-dir", type=Path, required=True)
    parser.add_argument("--blocks-dir", type=Path, required=True)
    parser.add_argument("--export-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    report = {
        "report_kind": "engineering_acceptance_not_accuracy",
        "database": str(args.db),
        "task1_coordinate_audit": audit_task1(args.task1_dir, args.blocks_dir),
        "task2_structure_and_coordinate_audit": audit_task2(args.relation_dir, args.blocks_dir),
        "api_verification": verify_api(args.db, args.blocks_dir, args.export_dir),
    }
    task1 = report["task1_coordinate_audit"]
    task2 = report["task2_structure_and_coordinate_audit"]
    failures: list[str] = []
    if task1["invalid_primary_references"] or task1["invalid_evidence_references"]:
        failures.append("task1 has invalid source coordinates")
    if task2["invalid_references"]:
        failures.append("task2 has invalid source coordinates")
    if task2["success_notices"] != task2["notices"]:
        failures.append("not every task2 notice succeeded")
    if not report["api_verification"]["passed"]:
        failures.extend(report["api_verification"]["failures"])
    report["failures"] = failures
    report["passed"] = not failures
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({
        "passed": report["passed"], "failures": failures,
        "task1_items": task1["items"], "task2_bids": task2["bids"],
        "api_timings_ms": report["api_verification"]["timings_ms"],
        "output": str(args.output),
    }, ensure_ascii=False, indent=2))
    return 0 if report["passed"] else 2


if __name__ == "__main__":
    raise SystemExit(main())
