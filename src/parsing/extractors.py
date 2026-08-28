"""Format-specific extractors that all emit the same :class:`Block` model."""

from __future__ import annotations

import json
import contextlib
import io
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from typing import Any, Callable, Iterable

from .schema import Block
from .utils import decode_bytes, normalise_text, table_to_markdown, trim_table


SourceFactory = Callable[..., dict[str, Any]]


def _block(block_type: str, text: str, source: dict[str, Any], **kwargs: Any) -> Block:
    return Block(type=block_type, text=normalise_text(text), source=source, **kwargs)


def _html_table_rows(table: Any) -> list[list[str]]:
    grid: list[list[str | None]] = []
    spans: dict[tuple[int, int], str] = {}
    tr_nodes = table.xpath(".//tr[not(ancestor::table[2])]")
    for row_index, tr in enumerate(tr_nodes):
        row: list[str | None] = []
        column = 0

        def fill_spans() -> None:
            nonlocal column
            while (row_index, column) in spans:
                while len(row) <= column:
                    row.append(None)
                row[column] = spans[(row_index, column)]
                column += 1

        fill_spans()
        for cell in tr.xpath("./th|./td"):
            fill_spans()
            value = normalise_text(cell.text_content())
            try:
                rowspan = max(1, int(cell.get("rowspan", "1")))
                colspan = max(1, int(cell.get("colspan", "1")))
            except ValueError:
                rowspan = colspan = 1
            for offset in range(colspan):
                while len(row) <= column + offset:
                    row.append(None)
                row[column + offset] = value
                for future_row in range(row_index + 1, row_index + rowspan):
                    spans[(future_row, column + offset)] = value
            column += colspan
        fill_spans()
        grid.append([cell or "" for cell in row])
    return trim_table(grid)


def parse_html(path: Path, source: SourceFactory) -> tuple[list[Block], list[str]]:
    warnings: list[str] = []
    try:
        from lxml import html as lxml_html
    except ImportError as exc:
        raise RuntimeError("HTML parsing requires lxml") from exc

    document = lxml_html.fromstring(decode_bytes(path.read_bytes()))
    for bad in document.xpath("//script|//style|//noscript|//svg|//nav|//footer"):
        bad.drop_tree()

    roots = document.xpath(
        "//*[contains(concat(' ', normalize-space(@class), ' '), ' vF_detail_main ')]"
    )
    if not roots:
        roots = document.xpath("//main|//article|//*[@id='content']|//*[@id='main-content']")
    root = roots[0] if roots else (document.xpath("//body") or [document])[0]

    # Generic pages often include site chrome inside body. Remove only strong layout signals.
    for bad in root.xpath(
        ".//*[self::aside or contains(translate(@class,'ABCDEFGHIJKLMNOPQRSTUVWXYZ','abcdefghijklmnopqrstuvwxyz'),'footer') "
        "or contains(translate(@class,'ABCDEFGHIJKLMNOPQRSTUVWXYZ','abcdefghijklmnopqrstuvwxyz'),'sidebar') "
        "or contains(translate(@class,'ABCDEFGHIJKLMNOPQRSTUVWXYZ','abcdefghijklmnopqrstuvwxyz'),'related')]"
    ):
        bad.drop_tree()

    blocks: list[Block] = []
    paragraph_index = 0
    table_index = 0
    block_tags = {"h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "dt", "dd", "pre"}
    for node in root.iterdescendants():
        tag = str(node.tag).lower() if isinstance(node.tag, str) else ""
        if node.xpath("ancestor::table"):
            continue
        if tag == "table":
            rows = _html_table_rows(node)
            if rows:
                table_index += 1
                blocks.append(
                    _block(
                        "table",
                        table_to_markdown(rows),
                        source(table_index=table_index),
                        rows=rows,
                        metadata={"row_count": len(rows), "column_count": max(map(len, rows))},
                    )
                )
            continue
        emit = tag in block_tags
        if tag in {"div", "section", "article"} and not node.xpath(
            ".//*[self::p or self::table or self::ul or self::ol or self::h1 or self::h2 or self::h3 or self::h4 or self::h5 or self::h6]"
        ):
            emit = True
        if emit:
            text = normalise_text(node.text_content())
            if text:
                paragraph_index += 1
                blocks.append(
                    _block(
                        "heading" if tag.startswith("h") else "paragraph",
                        text,
                        source(paragraph_index=paragraph_index),
                        metadata={"html_tag": tag},
                    )
                )
    if not blocks:
        text = normalise_text(root.text_content())
        if text:
            blocks.append(_block("paragraph", text, source(paragraph_index=1)))
        else:
            warnings.append("html_no_content")
    return blocks, warnings


class OCRBackend:
    """Lazy OCR adapter. RapidOCR is preferred; Tesseract is a fallback."""

    def __init__(self, mode: str = "auto") -> None:
        self.mode = mode
        self._engine: Any = None
        self._name: str | None = None
        self._initialised = False

    @property
    def name(self) -> str | None:
        self._initialise()
        return self._name

    def _initialise(self) -> None:
        if self._initialised:
            return
        self._initialised = True
        if self.mode == "none":
            return
        if self.mode in {"auto", "rapidocr"}:
            try:
                from rapidocr_onnxruntime import RapidOCR

                self._engine = RapidOCR()
                self._name = "rapidocr"
                return
            except Exception:
                if self.mode == "rapidocr":
                    return
        if self.mode in {"auto", "tesseract"}:
            try:
                import pytesseract

                if shutil.which("tesseract"):
                    self._engine = pytesseract
                    self._name = "tesseract"
            except Exception:
                return

    def extract(self, image_path: Path) -> tuple[str, dict[str, Any]]:
        self._initialise()
        if self._name == "rapidocr":
            result, _ = self._engine(str(image_path))
            lines: list[str] = []
            confidences: list[float] = []
            for item in result or []:
                if len(item) >= 3:
                    lines.append(normalise_text(item[1]))
                    try:
                        confidences.append(float(item[2]))
                    except (TypeError, ValueError):
                        pass
            return "\n".join(line for line in lines if line), {
                "ocr_engine": self._name,
                "mean_confidence": round(sum(confidences) / len(confidences), 4) if confidences else None,
            }
        if self._name == "tesseract":
            from PIL import Image, ImageOps

            image = ImageOps.exif_transpose(Image.open(image_path)).convert("RGB")
            return normalise_text(self._engine.image_to_string(image, lang="chi_sim+eng")), {
                "ocr_engine": self._name
            }
        return "", {"ocr_engine": None, "ocr_status": "unavailable"}


def parse_image(path: Path, source: SourceFactory, ocr: OCRBackend) -> tuple[list[Block], list[str]]:
    text, metadata = ocr.extract(path)
    warnings = [] if text else ["ocr_unavailable_or_empty"]
    block = _block("ocr_text", text, source(image_name=path.name), metadata=metadata)
    return [block], warnings


def _iter_docx_content(document: Any) -> Iterable[tuple[str, Any]]:
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield "paragraph", Paragraph(child, document)
        elif child.tag.endswith("}tbl"):
            yield "table", Table(child, document)


def parse_docx(
    path: Path,
    source: SourceFactory,
    ocr: OCRBackend,
    temp_dir: Path,
    min_embedded_image_pixels: int = 40_000,
) -> tuple[list[Block], list[str]]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("DOCX parsing requires python-docx") from exc

    document = Document(str(path))
    blocks: list[Block] = []
    warnings: list[str] = []
    paragraph_index = 0
    table_index = 0
    for kind, item in _iter_docx_content(document):
        if kind == "paragraph":
            text = normalise_text(item.text)
            if text:
                paragraph_index += 1
                style = normalise_text(getattr(getattr(item, "style", None), "name", ""))
                block_type = "heading" if style.lower().startswith("heading") or style.startswith("标题") else "paragraph"
                blocks.append(
                    _block(block_type, text, source(paragraph_index=paragraph_index), metadata={"style": style})
                )
        else:
            rows = trim_table([[cell.text for cell in row.cells] for row in item.rows])
            if rows:
                table_index += 1
                blocks.append(
                    _block(
                        "table",
                        table_to_markdown(rows),
                        source(table_index=table_index),
                        rows=rows,
                        metadata={"row_count": len(rows), "column_count": max(map(len, rows))},
                    )
                )

    # Embedded image order is difficult to infer reliably from all DOCX producers;
    # preserve the part name so it can still be audited and cited.
    with zipfile.ZipFile(path) as archive:
        media_names = sorted(name for name in archive.namelist() if name.startswith("word/media/") and not name.endswith("/"))
        for image_index, member in enumerate(media_names, 1):
            image_path = temp_dir / f"docx_image_{image_index}_{Path(member).name}"
            image_path.write_bytes(archive.read(member))
            image_pixels: int | None = None
            try:
                from PIL import Image

                with Image.open(image_path) as image:
                    image_pixels = image.width * image.height
            except Exception:
                pass
            if image_pixels is not None and image_pixels < min_embedded_image_pixels:
                text, metadata = "", {
                    "ocr_engine": ocr.name,
                    "ocr_status": "skipped_small_image",
                    "image_pixels": image_pixels,
                }
            else:
                try:
                    text, metadata = ocr.extract(image_path)
                except Exception as exc:
                    text, metadata = "", {
                        "ocr_engine": ocr.name,
                        "ocr_status": "unsupported_or_failed_image",
                        "error": f"{type(exc).__name__}: {exc}",
                    }
            metadata.update({"embedded_part": member})
            blocks.append(
                _block(
                    "ocr_text",
                    text,
                    source(image_name=member, image_index=image_index),
                    metadata=metadata,
                )
            )
            if not text and metadata.get("ocr_status") != "skipped_small_image":
                warnings.append(f"embedded_image_ocr_empty:{member}:{metadata.get('ocr_status', 'empty')}")
    if not blocks:
        warnings.append("docx_no_content")
    return blocks, warnings


def _sheet_blocks(
    sheet_name: str,
    rows: list[list[Any]],
    source: SourceFactory,
    *,
    max_rows: int = 500,
) -> list[Block]:
    clean = trim_table(rows)
    blocks: list[Block] = []
    for start in range(0, len(clean), max_rows):
        chunk = clean[start : start + max_rows]
        if not chunk:
            continue
        # Retain the first row as a likely header for later chunks.
        serialised = chunk if start == 0 else [clean[0], *chunk]
        blocks.append(
            _block(
                "table",
                table_to_markdown(serialised),
                source(sheet=sheet_name, row_start=start + 1, row_end=min(start + max_rows, len(clean))),
                rows=serialised,
                metadata={
                    "row_count": len(serialised),
                    "column_count": max(map(len, serialised)),
                    "header_repeated": start > 0,
                },
            )
        )
    return blocks


def parse_xlsx(path: Path, source: SourceFactory) -> tuple[list[Block], list[str]]:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("XLSX parsing requires openpyxl") from exc

    workbook = openpyxl.load_workbook(path, data_only=True, read_only=False)
    blocks: list[Block] = []
    warnings: list[str] = []
    try:
        for sheet in workbook.worksheets:
            merged_values: dict[tuple[int, int], Any] = {}
            for merged in sheet.merged_cells.ranges:
                anchor = sheet.cell(merged.min_row, merged.min_col).value
                for row in range(merged.min_row, merged.max_row + 1):
                    for column in range(merged.min_col, merged.max_col + 1):
                        merged_values[(row, column)] = anchor
            rows: list[list[Any]] = []
            for row in sheet.iter_rows():
                values = [merged_values.get((cell.row, cell.column), cell.value) for cell in row]
                rows.append(values)
            blocks.extend(_sheet_blocks(sheet.title, rows, source))
            if not trim_table(rows):
                warnings.append(f"empty_sheet:{sheet.title}")
    finally:
        workbook.close()
    return blocks, warnings


def parse_xls(path: Path, source: SourceFactory) -> tuple[list[Block], list[str]]:
    try:
        import xlrd
    except ImportError as exc:
        raise RuntimeError("XLS parsing requires xlrd") from exc

    workbook = xlrd.open_workbook(path, on_demand=True)
    blocks: list[Block] = []
    warnings: list[str] = []
    try:
        for sheet in workbook.sheets():
            rows = [[sheet.cell_value(r, c) for c in range(sheet.ncols)] for r in range(sheet.nrows)]
            for rlo, rhi, clo, chi in getattr(sheet, "merged_cells", []):
                anchor = rows[rlo][clo]
                for r in range(rlo, rhi):
                    for c in range(clo, chi):
                        rows[r][c] = anchor
            blocks.extend(_sheet_blocks(sheet.name, rows, source))
            if not trim_table(rows):
                warnings.append(f"empty_sheet:{sheet.name}")
    finally:
        workbook.release_resources()
    return blocks, warnings


def bbox_overlap_ratio(subject: Iterable[float], reference: Iterable[float]) -> float:
    """Return the fraction of ``subject`` covered by ``reference`` bboxes."""

    ax0, ay0, ax1, ay1 = (float(value) for value in subject)
    bx0, by0, bx1, by1 = (float(value) for value in reference)
    width = max(0.0, min(ax1, bx1) - max(ax0, bx0))
    height = max(0.0, min(ay1, by1) - max(ay0, by0))
    subject_area = max(0.0, ax1 - ax0) * max(0.0, ay1 - ay0)
    return (width * height / subject_area) if subject_area else 0.0


PDF_WATERMARK_TOKEN_RE = re.compile(r"^[0-9a-fA-F]{32}-\d{14,20}$")


def _normalised_rgb(value: Any) -> tuple[float, float, float] | None:
    """Normalise PyMuPDF / pdfplumber colours to an RGB tuple in ``[0, 1]``."""

    if isinstance(value, int):
        return (
            ((value >> 16) & 255) / 255,
            ((value >> 8) & 255) / 255,
            (value & 255) / 255,
        )
    if not isinstance(value, (list, tuple)) or len(value) < 3:
        return None
    try:
        rgb = tuple(float(component) for component in value[:3])
    except (TypeError, ValueError):
        return None
    if max(rgb) > 1:
        rgb = tuple(component / 255 for component in rgb)
    return rgb


def detect_pdf_watermarks(page: Any) -> list[dict[str, Any]]:
    """Detect only strong, page-native watermark tokens before table extraction.

    The identifier-shaped text alone is deliberately insufficient.  A candidate
    must also look like a watermark (rotated, translucent, or mid-grey).  This
    prevents legitimate horizontal black identifiers from being removed.
    """

    watermarks: list[dict[str, Any]] = []
    try:
        page_dict = page.get_text("dict")
    except Exception:
        return watermarks
    for block in page_dict.get("blocks", []):
        for line in block.get("lines", []):
            direction = line.get("dir", (1.0, 0.0))
            try:
                direction = (float(direction[0]), float(direction[1]))
            except (IndexError, TypeError, ValueError):
                direction = (1.0, 0.0)
            for span in line.get("spans", []):
                raw_text = str(span.get("text", ""))
                compact_text = re.sub(r"\s+", "", raw_text)
                if not PDF_WATERMARK_TOKEN_RE.fullmatch(compact_text):
                    continue
                rgb = _normalised_rgb(span.get("color"))
                rotated = abs(direction[1]) >= 0.15
                try:
                    alpha = float(span.get("alpha", 255))
                except (TypeError, ValueError):
                    alpha = 255
                translucent = alpha < 245
                mid_grey = bool(
                    rgb
                    and max(rgb) - min(rgb) <= 0.04
                    and 0.15 <= sum(rgb) / 3 <= 0.85
                )
                if not (rotated or translucent or mid_grey):
                    continue
                try:
                    bbox = tuple(float(value) for value in span["bbox"])
                    size = float(span.get("size", 0.0))
                except (KeyError, TypeError, ValueError):
                    continue
                watermarks.append(
                    {
                        "text": compact_text,
                        "raw_text": raw_text,
                        "bbox": bbox,
                        "font": str(span.get("font", "")),
                        "size": size,
                        "rgb": rgb,
                        "direction": direction,
                        "rotated": rotated,
                        "translucent": translucent,
                        "mid_grey": mid_grey,
                    }
                )
    return watermarks


def _strip_detected_pdf_watermarks(text: Any, watermarks: list[dict[str, Any]]) -> str:
    """Remove already-confirmed complete tokens from native paragraph text."""

    cleaned = str(text or "")
    for watermark in watermarks:
        raw_text = watermark["raw_text"]
        if raw_text:
            cleaned = cleaned.replace(raw_text, "")
        token = watermark["text"]
        if token != raw_text:
            cleaned = cleaned.replace(token, "")
    return normalise_text(cleaned)


def _pdf_char_matches_watermark(char: dict[str, Any], watermark: dict[str, Any]) -> bool:
    if str(char.get("fontname", "")) != watermark["font"]:
        return False
    try:
        size = float(char.get("size", 0.0))
    except (TypeError, ValueError):
        return False
    if abs(size - watermark["size"]) > max(1.0, watermark["size"] * 0.08):
        return False

    expected_rgb = watermark["rgb"]
    actual_rgb = _normalised_rgb(char.get("non_stroking_color"))
    if expected_rgb and (
        not actual_rgb
        or any(abs(actual - expected) > 0.04 for actual, expected in zip(actual_rgb, expected_rgb))
    ):
        return False

    if watermark["rotated"]:
        if char.get("upright") is True:
            return False
        matrix = char.get("matrix")
        if not isinstance(matrix, (list, tuple)) or len(matrix) < 2:
            return False
        try:
            char_direction = (float(matrix[0]), float(matrix[1]))
        except (TypeError, ValueError):
            return False
        expected_direction = watermark["direction"]
        if (
            abs(abs(char_direction[0]) - abs(expected_direction[0])) > 0.1
            or abs(abs(char_direction[1]) - abs(expected_direction[1])) > 0.1
        ):
            return False

    try:
        x0 = float(char["x0"])
        x1 = float(char["x1"])
        top = float(char["top"] if "top" in char else char["y0"])
        bottom = float(char["bottom"] if "bottom" in char else char["y1"])
    except (KeyError, TypeError, ValueError):
        return False
    left, bbox_top, right, bbox_bottom = watermark["bbox"]
    centre_x = (x0 + x1) / 2
    centre_y = (top + bottom) / 2
    return left - 1 <= centre_x <= right + 1 and bbox_top - 1 <= centre_y <= bbox_bottom + 1


def extract_pdf_table_with_watermark_filter(
    table: Any,
    watermarks: list[dict[str, Any]],
) -> tuple[list[list[Any]], int, list[str]]:
    """Extract a table after conservatively removing verified watermark chars.

    PyMuPDF's table finder copies page characters into ``table._chars``.  A
    rotated watermark is otherwise split across unrelated cells.  Characters
    are removed only when their complete ordered sequence reproduces a strong
    page-level watermark; partial matches are intentionally left untouched.
    """

    chars = getattr(table, "_chars", None)
    if not isinstance(chars, list) or not watermarks:
        return table.extract(), 0, []

    original_chars = list(chars)
    removed_indexes: set[int] = set()
    matched_patterns: list[str] = []
    for watermark in watermarks:
        indexes = [
            index
            for index, char in enumerate(original_chars)
            if index not in removed_indexes and _pdf_char_matches_watermark(char, watermark)
        ]
        candidate_text = "".join(str(original_chars[index].get("text", "")) for index in indexes)
        if candidate_text == watermark["text"]:
            removed_indexes.update(indexes)
            matched_patterns.append(watermark["text"])

    if not removed_indexes:
        return table.extract(), 0, []
    chars[:] = [char for index, char in enumerate(original_chars) if index not in removed_indexes]
    try:
        rows = table.extract()
    finally:
        chars[:] = original_chars
    return rows, len(removed_indexes), matched_patterns


def parse_pdf(
    path: Path,
    source: SourceFactory,
    ocr: OCRBackend,
    temp_dir: Path,
    scan_text_threshold: int = 50,
    table_text_overlap_threshold: float = 0.5,
    max_ocr_pages: int = 20,
    skip_ocr_if_scan_pages_over: int = 0,
    filter_native_watermarks: bool = True,
) -> tuple[list[Block], list[str]]:
    try:
        import pymupdf as fitz
    except ImportError as exc:
        raise RuntimeError("PDF parsing requires PyMuPDF") from exc

    document = fitz.open(path)
    blocks: list[Block] = []
    warnings: list[str] = []
    try:
        page_watermarks = [
            detect_pdf_watermarks(page) if filter_native_watermarks else []
            for page in document
        ]
        native_texts = [
            _strip_detected_pdf_watermarks(page.get_text(), page_watermarks[index])
            for index, page in enumerate(document)
        ]
        scan_pages = [index for index, text in enumerate(native_texts) if len(text) < scan_text_threshold]
        selected_scan_pages = set(scan_pages)
        ocr_available = ocr.name is not None
        if skip_ocr_if_scan_pages_over > 0 and len(scan_pages) > skip_ocr_if_scan_pages_over:
            selected_scan_pages = set()
            warnings.append(
                f"pdf_ocr_skipped_long_scan:scan_pages={len(scan_pages)}:threshold={skip_ocr_if_scan_pages_over}"
            )
        elif not ocr_available and scan_pages:
            selected_scan_pages = set()
            warnings.append(f"ocr_unavailable:scan_pages={len(scan_pages)}")
        elif max_ocr_pages > 0 and len(scan_pages) > max_ocr_pages:
            keywords = ("主要标的信息", "报价", "品牌", "型号", "单价", "数量", "中标", "成交")
            priority = [index for index in scan_pages if any(word in native_texts[index] for word in keywords)]
            chosen = priority[:max_ocr_pages]
            remaining_slots = max_ocr_pages - len(chosen)
            candidates = [index for index in scan_pages if index not in chosen]
            if remaining_slots > 0 and candidates:
                if remaining_slots == 1:
                    sampled = [candidates[len(candidates) // 2]]
                else:
                    sampled = [
                        candidates[round(position * (len(candidates) - 1) / (remaining_slots - 1))]
                        for position in range(remaining_slots)
                    ]
                chosen.extend(sampled)
            selected_scan_pages = set(chosen)
            warnings.append(f"pdf_ocr_page_limit:selected={len(selected_scan_pages)}:skipped={len(scan_pages) - len(selected_scan_pages)}")

        for page_index, page in enumerate(document, 1):
            watermarks = page_watermarks[page_index - 1]
            text_blocks = sorted(page.get_text("blocks"), key=lambda item: (round(item[1]), item[0]))
            page_text_length = sum(len(_strip_detected_pdf_watermarks(raw[4], watermarks)) for raw in text_blocks)
            detected_tables: list[tuple[list[list[str]], list[float], int, list[str]]] = []
            try:
                finder = getattr(page, "find_tables", None)
                with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                    tables = finder().tables if finder else []
                for table in tables:
                    extracted, watermark_char_count, watermark_patterns = (
                        extract_pdf_table_with_watermark_filter(table, watermarks)
                    )
                    rows = trim_table(extracted)
                    if rows:
                        detected_tables.append(
                            (
                                rows,
                                [float(value) for value in table.bbox],
                                watermark_char_count,
                                watermark_patterns,
                            )
                        )
            except Exception as exc:
                warnings.append(f"pdf_table_detection_failed:page={page_index}:{type(exc).__name__}")

            table_bboxes = [bbox for _, bbox, _, _ in detected_tables]
            paragraph_index = 0
            for raw in text_blocks:
                text = _strip_detected_pdf_watermarks(raw[4], watermarks)
                if not text:
                    continue
                overlaps_table = table_text_overlap_threshold > 0 and any(
                    bbox_overlap_ratio(raw[:4], table_bbox) >= table_text_overlap_threshold
                    for table_bbox in table_bboxes
                )
                if overlaps_table:
                    continue
                paragraph_index += 1
                blocks.append(
                    _block(
                        "paragraph",
                        text,
                        source(page=page_index, paragraph_index=paragraph_index),
                        metadata={"bbox": [round(float(v), 2) for v in raw[:4]]},
                    )
                )

            for table_index, (rows, table_bbox, watermark_char_count, watermark_patterns) in enumerate(
                detected_tables, 1
            ):
                metadata = {
                    "row_count": len(rows),
                    "column_count": max(map(len, rows)),
                    "bbox": [round(value, 2) for value in table_bbox],
                    "overlap_suppression_threshold": table_text_overlap_threshold,
                }
                if watermark_char_count:
                    metadata.update(
                        {
                            "watermark_filtered": True,
                            "watermark_char_count": watermark_char_count,
                            "watermark_patterns": watermark_patterns,
                        }
                    )
                blocks.append(
                    _block(
                        "table",
                        table_to_markdown(rows),
                        source(page=page_index, table_index=table_index),
                        rows=rows,
                        metadata=metadata,
                    )
                )

            zero_index = page_index - 1
            if page_text_length < scan_text_threshold and zero_index in selected_scan_pages:
                pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                image_path = temp_dir / f"pdf_page_{page_index}.png"
                pixmap.save(str(image_path))
                text, metadata = ocr.extract(image_path)
                metadata.update({"scan_page": True, "native_text_chars": page_text_length})
                blocks.append(_block("ocr_text", text, source(page=page_index), metadata=metadata))
                if not text:
                    warnings.append(f"pdf_scan_ocr_empty:page={page_index}")
            elif page_text_length < scan_text_threshold:
                blocks.append(
                    _block(
                        "ocr_text",
                        "",
                        source(page=page_index),
                        metadata={
                            "scan_page": True,
                            "native_text_chars": page_text_length,
                            "ocr_engine": ocr.name,
                            "ocr_status": "skipped_page_limit" if ocr_available else "unavailable",
                        },
                    )
                )
    finally:
        document.close()
    return blocks, warnings


def parse_json_or_text(path: Path, source: SourceFactory, file_type: str) -> tuple[list[Block], list[str]]:
    text = decode_bytes(path.read_bytes())
    if file_type == "json":
        try:
            payload = json.loads(text)
            text = json.dumps(payload, ensure_ascii=False, indent=2)
        except json.JSONDecodeError:
            return [_block("metadata", text, source())], ["invalid_json"]
        return [_block("metadata", text, source())], []
    return [_block("paragraph", text, source(paragraph_index=1))], []


def convert_legacy_office(
    path: Path,
    output_dir: Path,
    target: str,
    backend: str = "auto",
) -> tuple[Path | None, str | None, str | None]:
    """Convert binary Office files with LibreOffice or local MS Office COM."""

    candidate = output_dir / f"{path.stem}.{target}"
    failures: list[str] = []
    if backend == "none":
        return None, None, "legacy Office conversion is disabled"
    if backend in {"auto", "libreoffice"}:
        project_root = Path(__file__).resolve().parents[2]
        configured = os.environ.get("LIBREOFFICE_PATH")
        known_paths = [
            Path(configured) if configured else None,
            project_root / "third_party" / "LibreOffice" / "program" / "soffice.exe",
            Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
            Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"),
        ]
        executable = shutil.which("soffice") or shutil.which("libreoffice")
        executable = executable or next((str(item) for item in known_paths if item and item.exists()), None)
        if executable:
            profile = output_dir / f"lo_profile_{path.stem}"
            profile.mkdir(parents=True, exist_ok=True)
            command = [
                executable,
                f"-env:UserInstallation={profile.resolve().as_uri()}",
                "--headless",
                "--convert-to",
                target,
                "--outdir",
                str(output_dir),
                str(path),
            ]
            completed = subprocess.run(
                command,
                capture_output=True,
                text=True,
                encoding="gb18030" if os.name == "nt" else "utf-8",
                errors="ignore",
                timeout=180,
                check=False,
            )
            if completed.returncode == 0 and candidate.exists():
                return candidate, "libreoffice", None
            failures.append(f"LibreOffice failed: {completed.stderr.strip() or completed.stdout.strip()}")
        else:
            failures.append("LibreOffice executable was not found")
        if backend == "libreoffice":
            return None, None, "; ".join(failures)

    if backend in {"auto", "ms-office"} and os.name == "nt":
        powershell = shutil.which("powershell") or shutil.which("pwsh")
        script = Path(__file__).resolve().parents[2] / "scripts" / "office_convert.ps1"
        if powershell and script.exists():
            command = [
                powershell,
                "-NoProfile",
                "-NonInteractive",
                "-ExecutionPolicy",
                "Bypass",
                "-File",
                str(script),
                "-InputPath",
                str(path.resolve()),
                "-OutputPath",
                str(candidate.resolve()),
                "-Target",
                target,
            ]
            completed = subprocess.run(
                command,
                capture_output=True,
                text=True,
                encoding="gb18030",
                errors="ignore",
                timeout=180,
                check=False,
            )
            if completed.returncode == 0 and candidate.exists():
                return candidate, "ms-office", None
            failures.append(f"Microsoft Office conversion failed: {completed.stderr.strip() or completed.stdout.strip()}")
        else:
            failures.append("PowerShell converter was not found")
    return None, None, "; ".join(failures) or "no legacy Office converter is available"
