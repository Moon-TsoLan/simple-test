from __future__ import annotations

import csv
import io
import json
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from src.api.app import create_app


FIXTURE = Path(__file__).resolve().parents[1] / "fixtures" / "task2_synthetic.json"


def _client(tmp_path: Path) -> TestClient:
    return TestClient(create_app(tmp_path / "application.db", frontend_dist=tmp_path / "no-dist", upload_dir=tmp_path / "uploads"))


def test_health_empty_states_and_not_found(tmp_path: Path) -> None:
    with _client(tmp_path) as client:
        assert client.get("/api/health").json()["status"] == "ok"
        status = client.get("/api/status").json()
        assert status["task1"] == {"entities": 0, "notices": 0}
        assert status["task2"]["projects"] == 0
        assert status["model"]["status"] == "not_configured"
        assert client.get("/api/task1/items").json()["items"] == []
        assert client.get("/api/task1/items/missing").status_code == 404
        assert client.get("/api/jobs/missing").status_code == 404


def test_task1_import_search_detail_and_full_export(tmp_path: Path) -> None:
    notice_dir = tmp_path / "notices"
    notice_dir.mkdir()
    items = []
    for index in range(125):
        items.append({
            "item_no": index + 1,
            "product_service_name": f"测试产品{index + 1}",
            "category_name": "测试品目",
            "brand_supplier": "测试品牌",
            "unit_price": 10,
            "quantity": 2,
            "quantity_unit": "台",
            "total_price": 20,
            "source_file": "结果表.xlsx",
            "source_block_id": "N-001:blk",
            "source_block_row": index + 2,
            "evidence_refs": [{"block_id": "N-001:blk", "block_row": index + 2}],
        })
    (notice_dir / "N-001.json").write_text(
        json.dumps({"notice_id": "N-001", "title": "任务一合成公告", "items": items}, ensure_ascii=False),
        encoding="utf-8",
    )

    with _client(tmp_path) as client:
        report = client.post("/api/admin/import-task1", params={"path": str(notice_dir)}).json()
        assert report == {"notices": 1, "entities": 125, "failures": 0}
        page = client.get("/api/task1/items", params={"brand": "测试品牌", "page_size": 20}).json()
        assert page["total"] == 125 and len(page["items"]) == 20
        detail = client.get(f"/api/task1/items/{page['items'][0]['entity_id']}").json()
        assert detail["evidence_refs"][0]["block_id"] == "N-001:blk"

        csv_response = client.get("/api/task1/export.csv", params={"brand": "测试品牌"})
        csv_rows = list(csv.DictReader(io.StringIO(csv_response.content.decode("utf-8-sig"))))
        assert len(csv_rows) == 125

        xlsx_response = client.get("/api/task1/export.xlsx", params={"brand": "测试品牌"})
        workbook = load_workbook(io.BytesIO(xlsx_response.content), read_only=True)
        assert workbook.active.max_row == 126


def test_task2_endpoints_candidate_upload_and_graph(tmp_path: Path) -> None:
    projects = json.loads(FIXTURE.read_text(encoding="utf-8"))
    with _client(tmp_path) as client:
        for project in projects:
            assert client.post("/api/task2/projects", json=project).status_code == 200
        project_rows = client.get("/api/task2/projects").json()
        assert len(project_rows) == 3

        organizations = client.get("/api/task2/organizations").json()
        by_compact_name = {value["name"].replace(" ", ""): value["org_id"] for value in organizations}
        unit_id = by_compact_name["测试采购单位甲"]
        supplier_a = by_compact_name["甲供应商有限公司"]
        supplier_b = by_compact_name["乙供应商有限公司"]
        assert len(client.get(f"/api/task2/units/{unit_id}/win-suppliers").json()) == 2
        assert client.get(f"/api/task2/units/{unit_id}/top-bidders").json()[0]["bid_count"] == 2
        assert client.get(f"/api/task2/units/{unit_id}/co-bid-pairs").json()[0]["project_count"] == 2
        assert client.get(f"/api/task2/suppliers/{supplier_a}/co-bidders").status_code == 200
        ids = f"{supplier_a},{supplier_b}"
        assert len(client.get("/api/task2/suppliers/common-units", params={"ids": ids}).json()) == 2
        assert len(client.get("/api/task2/suppliers/joint-projects", params={"ids": ids}).json()) == 3
        graph = client.get("/api/task2/graph/subgraph", params={"project_id": project_rows[0]["project_id"]}).json()
        assert graph["nodes"] and graph["edges"]

        candidate = client.post("/api/task2/extract-candidate", json={
            "notice_id": "C-1", "title": "规则候选", "blocks": [{
                "type": "table", "block_id": "C-1:b1",
                "rows": [["供应商名称", "报价", "结果"], ["测试公司", "99万元", "中标"]],
            }],
        })
        assert candidate.status_code == 200
        assert candidate.json()["packages"][0]["bidders"][0]["amount_yuan"] == 990000

        upload = client.post(
            "/api/jobs/upload",
            data={"kind": "task1"},
            files=[("files", ("notice.html", b"<html>ok</html>", "text/html"))],
        )
        assert upload.status_code == 200
        job = client.get(f"/api/jobs/{upload.json()['job_id']}").json()
        assert job["status"] == "staged"
        assert job["files"][0]["status"] == "ready"


def test_api_validation_errors(tmp_path: Path) -> None:
    with _client(tmp_path) as client:
        assert client.get("/api/task1/items", params={"page_size": 101}).status_code == 422
        assert client.post("/api/jobs/upload", data={"kind": "wrong"}, files=[("files", ("a.html", b"x", "text/html"))]).status_code == 422
        duplicate = client.post(
            "/api/jobs/upload",
            data={"kind": "task1"},
            files=[
                ("files", ("a?.html", b"x", "text/html")),
                ("files", ("a*.html", b"y", "text/html")),
            ],
        )
        assert duplicate.status_code == 400
