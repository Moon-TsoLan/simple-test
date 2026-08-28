from __future__ import annotations

import json
import zipfile
from pathlib import Path

import openpyxl
import pytest

from src.parsing.extractors import (
    bbox_overlap_ratio,
    detect_pdf_watermarks,
    extract_pdf_table_with_watermark_filter,
    parse_docx,
)
from src.parsing.pipeline import NoticeParser, ParserConfig, group_archive_payloads
from src.parsing.utils import safe_archive_name


def test_html_and_xlsx_are_unified_and_traceable(tmp_path: Path) -> None:
    html_path = tmp_path / "N001.html"
    html_path.write_text(
        """
        <html><body><main>
          <h2>设备采购结果公告</h2><p>四、主要标的信息</p>
          <table><tr><th rowspan='2'>序号</th><th>货物名称</th></tr>
          <tr><td>台式计算机</td></tr></table>
        </main></body></html>
        """,
        encoding="utf-8",
    )
    workbook_path = tmp_path / "报价.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "报价明细"
    sheet.append(["产品名称", "品牌", "型号", "数量", "单价", "总价"])
    sheet.append(["台式计算机", "联想", "M450", 2, 5000, 10000])
    workbook.save(workbook_path)

    archive_path = tmp_path / "N001.zip"
    with zipfile.ZipFile(archive_path, "w") as archive:
        archive.write(workbook_path, "报价明细.xlsx")

    payload = NoticeParser(ParserConfig(ocr="none")).parse_notice("N001", html_path, archive_path)
    assert payload["schema_version"] == "1.1"
    assert payload["parser_version"] == "2.1.0"
    assert len(payload["parser_config_digest"]) == 64
    assert payload["title"] == "设备采购结果公告"
    assert all(block["block_id"].startswith("N001:blk_") for block in payload["blocks"])
    assert all(block["source"]["container_path"] for block in payload["blocks"])
    tables = [block for block in payload["blocks"] if block["type"] == "table"]
    assert any("台式计算机" in json.dumps(block["rows"], ensure_ascii=False) for block in tables)
    assert any(block["source"].get("sheet") == "报价明细" for block in tables)

    changed_config = NoticeParser(ParserConfig(ocr="none", max_ocr_pages_per_pdf=7)).parse_notice(
        "N001", html_path, archive_path
    )
    assert payload["parser_config_digest"] != changed_config["parser_config_digest"]
    assert {block["block_id"] for block in payload["blocks"]} == {
        block["block_id"] for block in changed_config["blocks"]
    }


def test_multipart_rar_volumes_are_collapsed_to_first_volume(tmp_path: Path) -> None:
    files = []
    for name in ("bundle.part3.rar", "readme.txt", "bundle.part1.rar", "bundle.part2.rar"):
        path = tmp_path / name
        path.write_bytes(b"fixture")
        files.append(path)
    payloads = group_archive_payloads(files, tmp_path)
    assert [relative for _, relative in payloads] == ["bundle.part1.rar", "readme.txt"]


def test_pdf_bbox_overlap_ratio_supports_table_text_suppression() -> None:
    assert bbox_overlap_ratio((10, 10, 20, 20), (0, 0, 30, 30)) == 1.0
    assert bbox_overlap_ratio((10, 10, 20, 20), (15, 10, 30, 20)) == 0.5
    assert bbox_overlap_ratio((10, 10, 20, 20), (30, 30, 40, 40)) == 0.0


def _watermark_span(
    text: str,
    *,
    direction=(0.5764, -0.8171),
    color=0x808080,
    alpha=128,
):
    class Page:
        @staticmethod
        def get_text(kind: str):
            assert kind == "dict"
            return {
                "blocks": [
                    {
                        "lines": [
                            {
                                "dir": direction,
                                "spans": [
                                    {
                                        "text": text,
                                        "font": "SimSun",
                                        "size": 16,
                                        "color": color,
                                        "alpha": alpha,
                                        "bbox": (100, 100, 500, 500),
                                    }
                                ],
                            }
                        ]
                    }
                ]
            }

    return Page()


def _pdf_char(text: str, *, upright: bool = False, color=(0.50196, 0.50196, 0.50196)):
    return {
        "text": text,
        "fontname": "SimSun",
        "size": 15.76,
        "non_stroking_color": color,
        "upright": upright,
        "matrix": (0.5764, 0.8171, -0.8171, 0.5764, 200, 200),
        "x0": 190,
        "x1": 210,
        "top": 190,
        "bottom": 210,
    }


def test_pdf_complete_rotated_watermark_is_removed_before_table_extraction() -> None:
    token = "5dc828f88bb244e488c521dda43bd6ca-20241211225211707"
    watermarks = detect_pdf_watermarks(_watermark_span(token))

    class Table:
        def __init__(self):
            self._chars = [_pdf_char(character) for character in token]
            self._chars.extend(
                _pdf_char(character, upright=True, color=(0.0, 0.0, 0.0))
                for character in "型号：拓之迹-JSGS01 Q235 DS-2CD3T26WDA4-L"
            )

        def extract(self):
            return [["".join(char["text"] for char in self._chars)]]

    table = Table()
    rows, removed, patterns = extract_pdf_table_with_watermark_filter(table, watermarks)
    assert rows == [["型号：拓之迹-JSGS01 Q235 DS-2CD3T26WDA4-L"]]
    assert removed == len(token)
    assert patterns == [token]
    assert len(table._chars) == len(token) + len(rows[0][0])


def test_pdf_partial_or_plain_identifier_is_not_removed() -> None:
    token = "5dc828f88bb244e488c521dda43bd6ca-20241211225211707"
    watermarks = detect_pdf_watermarks(_watermark_span(token))

    class PartialTable:
        def __init__(self):
            self._chars = [_pdf_char(character) for character in token[:-1]]

        def extract(self):
            return [["".join(char["text"] for char in self._chars)]]

    rows, removed, patterns = extract_pdf_table_with_watermark_filter(PartialTable(), watermarks)
    assert rows == [[token[:-1]]]
    assert removed == 0
    assert patterns == []
    assert detect_pdf_watermarks(
        _watermark_span(token, direction=(1, 0), color=0x000000, alpha=255)
    ) == []
    assert detect_pdf_watermarks(_watermark_span("型号：拓之迹-JSGS01")) == []


def test_unsupported_docx_image_is_isolated_instead_of_failing_document(tmp_path: Path) -> None:
    from docx import Document

    path = tmp_path / "embedded.docx"
    document = Document()
    document.add_paragraph("正文仍应保留")
    document.save(path)
    with zipfile.ZipFile(path, "a") as archive:
        archive.writestr("word/media/image1.wdp", b"unsupported-wdp-fixture")

    class FailingOCR:
        name = "rapidocr"

        @staticmethod
        def extract(_path: Path):
            raise ValueError("unsupported image")

    source = lambda **location: {"document_id": "f1", "file_name": path.name,
                                 "container_path": path.name, "file_type": "docx", **location}
    blocks, warnings = parse_docx(path, source, FailingOCR(), tmp_path)
    assert any(block.type == "paragraph" and block.text == "正文仍应保留" for block in blocks)
    image = next(block for block in blocks if block.type == "ocr_text")
    assert image.metadata["ocr_status"] == "unsupported_or_failed_image"
    assert any("image1.wdp" in warning for warning in warnings)


@pytest.mark.parametrize("name", ["../escape.txt", "/absolute.txt", "C:/escape.txt", "a/../../b.txt"])
def test_unsafe_archive_names_are_rejected(name: str) -> None:
    with pytest.raises(ValueError):
        safe_archive_name(name)


def test_one_broken_attachment_does_not_drop_other_files(tmp_path: Path) -> None:
    html_path = tmp_path / "N002.html"
    html_path.write_text("<html><body><main><p>中标公告正文</p></main></body></html>", encoding="utf-8")
    archive_path = tmp_path / "N002.zip"
    with zipfile.ZipFile(archive_path, "w") as archive:
        archive.writestr("说明.txt", "总价：100元")
        archive.writestr("损坏.pdf", b"%PDF-broken")

    payload = NoticeParser(ParserConfig(ocr="none")).parse_notice("N002", html_path, archive_path)
    assert any(block["text"] == "总价：100元" for block in payload["blocks"])
    assert any(block["type"] == "error" for block in payload["blocks"])
    assert payload["stats"]["document_status_counts"]["error"] == 1


def test_legacy_doc_fails_closed_without_emitting_binary_noise(tmp_path: Path) -> None:
    html_path = tmp_path / "N003.html"
    html_path.write_text("<html><body><main><p>公告正文</p></main></body></html>", encoding="utf-8")
    archive_path = tmp_path / "N003.zip"
    ole_header = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"Root Entry" * 100
    with zipfile.ZipFile(archive_path, "w") as archive:
        archive.writestr("legacy.doc", ole_header)

    config = ParserConfig(ocr="none", legacy_office_backend="none")
    payload = NoticeParser(config).parse_notice("N003", html_path, archive_path)
    doc_blocks = [block for block in payload["blocks"] if block["source"]["file_type"] == "doc"]
    assert [block["type"] for block in doc_blocks] == ["error"]
    assert "Root Entry" not in json.dumps(doc_blocks, ensure_ascii=False)
